#!/usr/bin/env python3
"""Relatorio v2: panorama da base inteira + fichas completas das construtoras
pesquisadas + lista completa de oportunidades + planilha.

Entradas: dados/grupos.json (05_mapa_base.py), dados/selecao_v2_gids.json,
dados/v2/indice.json, dados/v2/pesquisa/<slug>.json (pesquisa web profunda) e,
como reserva, dados/pesquisa/<dominio>.json (primeira rodada).
Saidas: fichas_v2/*.html, fichas_v2/relatorio.html e Construtoras_RMBH_mapeamento.xlsx."""
import collections, json, math, re, sys
from datetime import date
from html import escape
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import regionais_bh as RB
from regionais_bh import norm

R = Path(__file__).resolve().parent.parent
OUT = R / "fichas_v2"
HOJE = date.today().strftime("%d/%m/%Y")
G = json.load(open(R / "dados/grupos.json"))
GID = {g["gid"]: g for g in G}
SEL = json.load(open(R / "dados/selecao_v2_gids.json"))
SLUGS = json.load(open(R / "dados/v2/indice.json"))
RB.preparar([e["bairro"] for g in G for e in g["empreendimentos"]] + [g["bairro_sede"] for g in G])

BAIRROS_MEDIO = ["Serra", "Sion", "Buritis", "Palmeiras", "Estoril", "Ouro Preto", "Paquetá", "Santa Terezinha",
                 "Manacás", "Engenho Nogueira", "Grajaú", "Nova Suíça", "Gameleira"]
BAIRROS_LESTE = ["Alto Vera Cruz", "Boa Vista", "Pompéia", "São Geraldo", "Esplanada", "Taquaril", "Saudade"]
SINONIMOS = {"nova suica": {"nova suica", "nova suissa"}, "buritis": {"buritis", "dos buritis", "burutis"},
             "paqueta": {"paqueta", "jardim paqueta"}, "gameleira": {"gameleira", "nova gameleira"},
             "sion": {"sion", "carmo sion"}, "taquaril": {"taquaril", "conjunto taquaril"}}

# posicoes aproximadas (lat, lon) para o mapa esquematico
POS = {"Venda Nova": (-19.815, -43.960), "Norte": (-19.825, -43.915), "Pampulha": (-19.855, -43.985),
       "Nordeste": (-19.870, -43.905), "Noroeste": (-19.905, -43.985), "Leste": (-19.905, -43.895),
       "Centro-Sul": (-19.940, -43.935), "Oeste": (-19.955, -43.980), "Barreiro": (-19.985, -44.025),
       "Contagem": (-19.915, -44.070), "Betim": (-19.965, -44.185), "Ribeirão das Neves": (-19.765, -44.085),
       "Nova Lima": (-20.000, -43.850), "Sabará": (-19.885, -43.805), "Santa Luzia": (-19.770, -43.850),
       "Vespasiano": (-19.690, -43.920), "Lagoa Santa": (-19.630, -43.890), "Ibirité": (-20.020, -44.060),
       "Pedro Leopoldo": (-19.620, -44.045), "Sarzedo": (-20.035, -44.145), "Brumadinho": (-20.140, -44.200),
       "Esmeraldas": (-19.760, -44.310), "Sete Lagoas": (-19.530, -44.250), "Confins": (-19.630, -43.990),
       "Raposos": (-19.965, -43.805), "Mário Campos": (-20.055, -44.185), "São José da Lapa": (-19.700, -43.960),
       "Matozinhos": (-19.555, -44.080), "Igarapé": (-20.070, -44.300), "Juatuba": (-19.950, -44.340)}
CIDADES = {norm(k): k for k in POS if k not in RB._R}
REGIONAIS = list(RB._R)

# ------------------------------------------------------------------ util
def e(x):
    return escape(str(x)) if x not in (None, "", [], {}) else ""

def g_(d, *ks, default=""):
    for k in ks:
        if not isinstance(d, dict): return default
        d = d.get(k)
    return default if d in (None, "", [], {}) else d

def link(url, txt=None):
    if not url or not isinstance(url, str): return ""
    u = url if url.startswith("http") else "https://" + url
    t = txt or re.sub(r"^https?://(www\.)?", "", url).rstrip("/")
    return f'<a href="{e(u)}">{e(t)}</a>'

def fonte(url):
    return f' <a class="src" href="{e(url)}">fonte</a>' if isinstance(url, str) and url.startswith("http") else ""

def num(x, default=0.0):
    m = re.search(r"\d+(?:[.,]\d+)?", str(x if x is not None else ""))
    return float(m.group().replace(",", ".")) if m else default

def fmt_cnpj(c):
    c = re.sub(r"\D", "", str(c))
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}" if len(c) == 14 else c

def dig(t):
    return re.sub(r"\D", "", str(t or ""))

def brl(v):
    try: return "R$ " + f"{float(v):,.0f}".replace(",", ".")
    except (TypeError, ValueError): return ""

def local_de(bairro, cidade):
    """Onde fica: regional de BH (pelo bairro) ou a cidade da RMBH."""
    c = norm(cidade)
    if c and c not in ("belo horizonte", "bh"):
        return CIDADES.get(c, cidade.strip().title())
    reg = RB.regional(bairro) if bairro else "BH (sem regional)"
    if reg.startswith("BH") and bairro:
        # "Solar do Barreiro (Vale do Jatoba)" -> procura um bairro conhecido dentro do texto
        t = " " + re.sub(r"[^a-z ]", " ", norm(bairro)) + " "
        for nome in sorted(RB._MAPA, key=len, reverse=True):
            if len(nome) > 4 and f" {nome} " in t:
                return RB._MAPA[nome]
    return reg

def tabela(cab, linhas, cls=""):
    if not linhas: return '<p class="vazio">Nada encontrado.</p>'
    h = "".join(f"<th>{c}</th>" for c in cab)
    b = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in l) + "</tr>" for l in linhas)
    return f'<div class="wrap"><table class="{cls}"><thead><tr>{h}</tr></thead><tbody>{b}</tbody></table></div>'

# ------------------------------------------------------------------ graficos (SVG, uma serie, cor de destaque)
def barras(dados, larg=640, alt=180, fmt=str):
    """dados: lista de (rotulo, valor). Barras verticais, uma serie, tooltip nativo."""
    if not dados: return ""
    n = len(dados); m = max(v for _, v in dados) or 1
    pad_l, pad_b, pad_t = 28, 24, 16
    w = (larg - pad_l) / n; bw = max(6, w * 0.62)
    ticks = [0, m / 2, m]
    out = [f'<svg class="chart" viewBox="0 0 {larg} {alt}" role="img" aria-label="gráfico de barras">']
    for t in ticks:
        y = alt - pad_b - (alt - pad_b - pad_t) * t / m
        out.append(f'<line class="grid" x1="{pad_l}" x2="{larg}" y1="{y:.1f}" y2="{y:.1f}"/>'
                   f'<text class="ax" x="{pad_l - 4}" y="{y + 4:.1f}" text-anchor="end">{int(round(t))}</text>')
    for i, (rot, v) in enumerate(dados):
        h = (alt - pad_b - pad_t) * v / m
        x = pad_l + i * w + (w - bw) / 2; y = alt - pad_b - h
        r = min(4, bw / 2, h / 2) if h > 0 else 0
        out.append(f'<g class="bar"><title>{e(rot)}: {fmt(v)}</title>'
                   f'<path d="M{x:.1f},{alt - pad_b} V{y + r:.1f} Q{x:.1f},{y:.1f} {x + r:.1f},{y:.1f} H{x + bw - r:.1f} '
                   f'Q{x + bw:.1f},{y:.1f} {x + bw:.1f},{y + r:.1f} V{alt - pad_b} Z"/>'
                   f'<rect class="hit" x="{pad_l + i * w:.1f}" y="{pad_t}" width="{w:.1f}" height="{alt - pad_b - pad_t}"/></g>')
        if n <= 14 or i % 2 == 0:
            out.append(f'<text class="ax" x="{x + bw / 2:.1f}" y="{alt - 8}" text-anchor="middle">{e(rot)}</text>')
    if dados:
        i, (rot, v) = max(enumerate(dados), key=lambda t: t[1][1])
        x = pad_l + i * w + w / 2; y = alt - pad_b - (alt - pad_b - pad_t) * v / m
        out.append(f'<text class="lbl" x="{x:.1f}" y="{y - 4:.1f}" text-anchor="middle">{fmt(v)}</text>')
    out.append("</svg>")
    return "".join(out)

def barras_h(dados, larg=640, fmt=str):
    if not dados: return ""
    m = max(v for _, v in dados) or 1; lh = 22; alt = lh * len(dados) + 6; pl = 150
    out = [f'<svg class="chart" viewBox="0 0 {larg} {alt}" role="img" aria-label="gráfico de barras horizontais">']
    for i, (rot, v) in enumerate(dados):
        y = 3 + i * lh; w = (larg - pl - 50) * v / m
        out.append(f'<g class="bar"><title>{e(rot)}: {fmt(v)}</title><text class="ax" x="{pl - 6}" y="{y + 14}" text-anchor="end">{e(rot)}</text>'
                   f'<rect x="{pl}" y="{y + 3}" width="{max(w, 1):.1f}" height="{lh - 8}" rx="3"/>'
                   f'<text class="lbl" x="{pl + w + 5:.1f}" y="{y + 14}">{fmt(v)}</text></g>')
    out.append("</svg>")
    return "".join(out)

def _painel(contagem, nomes, caixa, larg, alt, m, titulo, sempre=()):
    lat0, lat1, lon0, lon1 = caixa
    def xy(p):
        la, lo = p
        return (30 + (lo - lon0) / (lon1 - lon0) * (larg - 60), 22 + (la - lat1) / (lat0 - lat1) * (alt - 50))
    out = [f'<svg class="chart mapa" viewBox="0 0 {larg} {alt}" role="img" aria-label="{e(titulo)}">',
           f'<text class="ax" x="6" y="14">{e(titulo)}</text>']
    for nome in nomes:
        v = contagem.get(nome, 0)
        if not v and nome not in sempre: continue
        x, y = xy(POS[nome]); r = 4 + 24 * math.sqrt(v / m) if v else 3
        out.append(f'<g class="pt{" on" if v else ""}"><title>{e(nome)}: {v}</title><circle cx="{x:.1f}" cy="{y:.1f}" r="{r:.1f}"/>'
                   f'<text class="ax" x="{x:.1f}" y="{y + r + 12:.1f}" text-anchor="middle">{e(nome)}</text>'
                   + (f'<text class="lbl" x="{x:.1f}" y="{y + 4:.1f}" text-anchor="middle">{v}</text>' if v else "") + "</g>")
    out.append("</svg>")
    return "".join(out)

def mapa(contagem, larg=640, alt=470, destaque=None, legenda=True):
    """Mapa esquematico em dois paineis: BH por regional (ampliado) e RMBH por
    cidade (BH somada num circulo so). Area do circulo proporcional a contagem."""
    bh = [n for n in POS if n in RB._R]
    rm = [n for n in POS if n not in RB._R]
    c_rm = dict(contagem); c_rm["BH"] = sum(contagem.get(n, 0) for n in bh)
    POS["BH"] = (-19.915, -43.940)
    m_bh = max([contagem.get(n, 0) for n in bh] + [1]); m_rm = max([c_rm.get(n, 0) for n in rm + ["BH"]] + [1])
    esq = _painel(contagem, bh, (-20.01, -19.80, -44.06, -43.87), 360, 330, m_bh, "Belo Horizonte, por regional", sempre=bh)
    dir_ = _painel(c_rm, ["BH"] + rm, (-20.16, -19.50, -44.36, -43.76), 360, 330, m_rm, "Região metropolitana, por cidade",
                   sempre=("BH", "Contagem", "Betim", "Ribeirão das Neves", "Nova Lima"))
    return f'<div class="mapas">{esq}{dir_}</div><p class="nota">Esquemático; posições aproximadas.</p>'

# ------------------------------------------------------------------ dados das fichas
def carregar(slug, g):
    p = R / "dados/v2/pesquisa" / f"{slug}.json"
    if p.exists():
        try: return json.load(open(p)), "v2"
        except json.JSONDecodeError: pass
    p = R / "dados/pesquisa" / f"{g['dominio']}.json"
    if g["dominio"] and p.exists():
        return converter_v1(json.load(open(p))), "v1"
    return {}, ""

def converter_v1(p):
    """Leva a pesquisa da primeira rodada para o formato v2 (reserva)."""
    return {"marca": p.get("nome_comercial"), "site": {"url": g_(p, "site", "url"), "confirmado": g_(p, "site", "confirmado_da_empresa") is True, "fonte": g_(p, "site", "fonte")},
            "redes": {"instagram": p.get("instagram") or {}, "linkedin": p.get("linkedin_empresa") or {}},
            "sede": {"endereco": g_(p, "endereco_comercial", "texto"), "fonte": g_(p, "endereco_comercial", "fonte")},
            "emails": [{"email": x.get("email"), "area": x.get("origem"), "confirmado_publicacao": x.get("validado") == "confirmado", "fonte": x.get("fonte")} for x in p.get("emails", [])],
            "telefones": [{"numero": x.get("numero"), "tipo": x.get("tipo"), "area": x.get("origem"), "confirmado_publicacao": x.get("validado") == "confirmado", "fonte": x.get("fonte")} for x in p.get("telefones", [])],
            "pessoas": [{"nome": d.get("nome"), "cargo": d.get("cargo"), "linkedin": d.get("linkedin"), "fonte": d.get("fonte")} for d in p.get("diretores_e_decisores", [])],
            "controle": {"controladores_finais": p.get("controladores_finais", [])},
            "empreendimentos": p.get("empreendimentos", []), "saude": {"reclame_aqui": g_(p, "saude", "reclame_aqui") or {},
            "recuperacao_judicial": g_(p, "saude", "recuperacao_judicial_ou_falencia")},
            "aderencia": {"nota_0a10": g_(p, "aderencia_hadron", "nota_0a10"), "encaixe": g_(p, "aderencia_hadron", "encaixe"),
                          "justificativa": g_(p, "aderencia_hadron", "justificativa")},
            "alertas": p.get("alertas", []), "confianca_0a10": p.get("confianca_geral_0a10")}

def contatos(g, p):
    """Junta e-mails e telefones da web e da Receita, sem repetir."""
    em, descartes = {}, collections.Counter()
    for x in p.get("emails") or []:
        k = str(x.get("email", "")).lower().strip()
        if "@" not in k: continue
        em[k] = dict(email=k, area=x.get("area") or "", pessoa=x.get("pessoa") or "",
                     status="confirmado" if x.get("confirmado_publicacao") else "publicado?", origem="web", fonte=x.get("fonte"))
    for x in g["emails"]:
        k = x["email"]
        if x["classe"] == "provedor_generico": descartes["e-mail genérico (pode ser pessoal)"] += 1; continue
        if k in em:
            em[k]["origem"] = "web + Receita"; continue
        em[k] = dict(email=k, area="Receita", pessoa="", origem=f"Receita ({x['n_cnpjs']} CNPJ)",
                     status={"contador": "contador", "outro_dominio": "domínio de terceiro"}.get(x["classe"], "cadastro Receita"), fonte="")
    fo = {}
    for x in p.get("telefones") or []:
        d = dig(x.get("numero"))
        if d.startswith("55") and len(d) > 11: d = d[2:]
        if len(d) < 10: continue
        fo[d[-8:]] = dict(numero=x.get("numero"), tipo=x.get("tipo") or "", area=x.get("area") or "",
                         status="confirmado" if x.get("confirmado_publicacao") else "publicado?", origem="web", fonte=x.get("fonte"))
    for x in g["fones"]:
        d = dig(x["numero"])
        if d[-8:] in fo:
            fo[d[-8:]]["origem"] = "web + Receita"; continue
        if x["tipo"] == "movel": descartes["celular da Receita (pode ser pessoal)"] += 1; continue
        fo[d[-8:]] = dict(numero=x["numero"], tipo=x["tipo"], area="Receita", origem=f"Receita ({x['n_cnpjs']} CNPJ)",
                          status="contador" if x["classe"] == "contador" else "cadastro Receita", fonte="")
    ordem = {"confirmado": 0, "publicado?": 1, "cadastro Receita": 2, "domínio de terceiro": 3, "contador": 4}
    return (sorted(em.values(), key=lambda x: ordem.get(x["status"], 5)),
            sorted(fo.values(), key=lambda x: ordem.get(x["status"], 5)), descartes)

def empreendimentos(g, p):
    web = []
    for x in p.get("empreendimentos") or []:
        if not x.get("nome"): continue
        x = dict(x); x["onde"] = local_de(x.get("bairro", ""), x.get("cidade", "")); web.append(x)
    nomes_web = {norm(x["nome"]) for x in web}
    base = [x for x in g["empreendimentos"] if not any(n and n in norm(x["razao"]) for n in nomes_web if len(n) > 4)]
    return web, base

def onde_atua(web, base):
    c = collections.Counter(x["onde"] for x in web)
    for x in base:
        if not x["mesmo_bairro_da_sede"]: c[x["regional"]] += 1
    c.pop("BH (sem regional)", None)
    return c

def validadores(g, p, em, fo, web):
    v = []
    add = lambda n, ok, det="": v.append((n, ok, det))
    add("Site oficial confirmado", "ok" if g_(p, "site", "confirmado") is True else "falha" if g_(p, "site", "confirmado") is False else "nd",
        link(g_(p, "site", "url") or g["dominio"]))
    add("Instagram da empresa", "ok" if g_(p, "redes", "instagram", "url") else "falha", link(g_(p, "redes", "instagram", "url")))
    add("LinkedIn da empresa", "ok" if g_(p, "redes", "linkedin", "url") else "falha", link(g_(p, "redes", "linkedin", "url")))
    conf = [x for x in em if x["status"] == "confirmado"]
    add("E-mail corporativo publicado pela empresa", "ok" if conf else "parcial" if any(x["status"] == "cadastro Receita" for x in em) else "falha",
        f"{len(conf)} confirmado(s) · {len(em)} no total")
    nn = [x for x in em if re.search(r"novos|terreno|negocio|incorpora|comercial|diretoria", (x["area"] or "") + x["email"])]
    add("E-mail de novos negócios/terrenos/comercial", "ok" if nn else "falha", ", ".join(e(x["email"]) for x in nn[:2]))
    fx = [x for x in fo if x["status"] == "confirmado" and x["tipo"] in ("fixo", "0800")]
    add("Telefone fixo publicado pela empresa", "ok" if fx else "parcial" if fo else "falha", e(fx[0]["numero"]) if fx else "")
    pes = p.get("pessoas") or []
    dec = [x for x in pes if x.get("papel") in ("incorporacao", "novos_negocios", "terrenos", "produto", "projetos")]
    add("Decisor de incorporação/novos negócios identificado", "ok" if dec else "parcial" if pes else "falha",
        ", ".join(e(x.get("nome")) for x in dec[:2]))
    add("Decisor com LinkedIn", "ok" if any(x.get("linkedin") for x in pes) else "falha",
        f"{sum(1 for x in pes if x.get('linkedin'))} de {len(pes)} pessoas")
    add("Controlador final identificado", "ok" if g_(p, "controle", "controladores_finais") else "parcial" if g["socios_pf"] else "falha", "")
    loc = [x for x in web if x.get("bairro") or (x.get("cidade") and norm(x["cidade"]) not in ("belo horizonte", ""))]
    add("Portfólio com local (3+ empreendimentos)", "ok" if len(loc) >= 3 else "parcial" if loc else "falha", f"{len(loc)} com local na web · {g['n_spe']} SPEs na Receita")
    lanc = [x for x in web if x.get("status") in ("lancamento", "em_obras", "previsto")]
    add("Lançando ou em obras agora", "ok" if lanc or g["spe_desde_2023"] else "falha", f"{len(lanc)} na web · {g['spe_desde_2023']} SPEs desde 2023")
    alvo = set(g["bairros_alvo_medio"] + g["bairros_alvo_leste"]) | {norm(x.get("bairro", "")) for x in web if norm(x.get("bairro", "")) in
            {norm(b) for b in BAIRROS_MEDIO + BAIRROS_LESTE}}
    add("Atua em bairro-alvo", "ok" if alvo else "falha", e(", ".join(sorted(alvo))[:80]))
    rj = g_(p, "saude", "recuperacao_judicial")
    add("Sem recuperação judicial/falência", "falha" if rj == "sim" else "ok" if rj else "nd", "")
    ra = g_(p, "saude", "reclame_aqui") or {}
    nota = num(ra.get("nota"), None) if ra.get("nota") else None
    add("Reclame Aqui", "nd" if nota is None else "ok" if nota >= 7 else "parcial" if nota >= 5 else "falha",
        " ".join(x for x in (str(ra.get("nota") or ""), str(ra.get("reputacao") or "")) if x) or "sem página")
    enc = g_(p, "aderencia", "encaixe")
    add("Encaixe no perfil Hadron", "ok" if enc in ("mcmv", "medio", "misto") else "falha" if enc in ("alto", "fora") else "nd", e(enc))
    return v

def pct(v):
    pts = {"ok": 1, "parcial": 0.5, "falha": 0, "nd": 0}
    return round(100 * sum(pts[s] for _, s, _ in v) / len(v))

ICONE = {"ok": "✔", "parcial": "◐", "falha": "✘", "nd": "?"}
ROT = {"ok": "atende", "parcial": "parcial", "falha": "não atende", "nd": "sem dado"}

# ------------------------------------------------------------------ ficha
def ficha(item, pos):
    g, p, slug, fonte_p = item["g"], item["p"], item["slug"], item["fonte"]
    em, fo, desc = item["em"], item["fo"], item["desc"]
    web, base, v = item["web"], item["base"], item["v"]
    marca = g_(p, "marca") or g["nome"]
    ad = p.get("aderencia") or {}
    site = g_(p, "site", "url") or g["dominio"]
    atua = item["atua"]

    em_l = [[f'<span class="mono">{e(x["email"])}</span>', e(x["area"]), e(x["pessoa"]), e(x["origem"]),
             f'<span class="st {x["status"].replace(" ", "_").replace("?", "")}">{e(x["status"])}</span>', fonte(x["fonte"])] for x in em]
    fo_l = [[f'<span class="mono">{e(x["numero"])}</span>', e(x["tipo"]), e(x["area"]), e(x["origem"]),
             f'<span class="st {x["status"].replace(" ", "_").replace("?", "")}">{e(x["status"])}</span>', fonte(x["fonte"])] for x in fo]
    pes = sorted(p.get("pessoas") or [], key=lambda x: ["incorporacao", "novos_negocios", "terrenos", "produto", "projetos",
                 "diretoria", "socio", "engenharia", "comercial", "marketing", "outro"].index(x.get("papel")) if x.get("papel") in
                 ["incorporacao", "novos_negocios", "terrenos", "produto", "projetos", "diretoria", "socio", "engenharia", "comercial", "marketing", "outro"] else 99)
    pes_l = [[f'<b>{e(x.get("nome"))}</b>', e(x.get("cargo")), e((x.get("papel") or "").replace("_", " ")),
              link(x.get("linkedin"), "LinkedIn") if x.get("linkedin") else "", e(x.get("formacao_ou_historico")), fonte(x.get("fonte"))] for x in pes]
    web_l = [[f'<b>{e(x.get("nome"))}</b>', e(", ".join(y for y in (x.get("endereco"), x.get("bairro"), x.get("cidade")) if y)),
              e(x.get("onde")), e((x.get("status") or "").replace("_", " ")), e(x.get("ano")), e(x.get("unidades")),
              e(x.get("area_privativa")), e(x.get("preco_m2") or x.get("preco")) + (" *" if x.get("calculado") else ""),
              e(x.get("padrao")), fonte(x.get("fonte"))] for x in web]
    base_l = [[f'<span class="mono">{fmt_cnpj(x["cnpj"])}</span>', e(x["razao"]), e(x["bairro"]) + (' <span class="muted">(= sede)</span>' if x["mesmo_bairro_da_sede"] else ""),
               e(x["regional"]), e(x["ano"])] for x in base[:30]]
    ctrl = g_(p, "controle") or {}
    cf = ", ".join(e(c.get("nome")) + (f' <span class="muted">({e(c.get("via"))})</span>' if c.get("via") else "") for c in ctrl.get("controladores_finais") or [])
    sa = p.get("saude") or {}
    ra = sa.get("reclame_aqui") or {}
    notic = "".join(f'<li><span class="tom {e(n.get("tom"))}">{e(n.get("tom") or "")}</span> {e(n.get("data"))} {e(n.get("resumo"))}{fonte(n.get("fonte"))}</li>'
                    for n in (sa.get("noticias") or []))
    proc = "".join(f'<li>{e(n.get("resumo"))}{fonte(n.get("fonte"))}</li>' for n in (sa.get("processos") or []))
    alertas = "".join(f"<li>{e(a)}</li>" for a in p.get("alertas") or [])
    canal = p.get("canal_para_oferecer_projeto") or {}
    at = p.get("atuacao") or {}
    redes = p.get("redes") or {}
    outros_end = "".join(f'<li>{e(o.get("tipo"))}: {e(o.get("endereco"))}{fonte(o.get("fonte"))}</li>' for o in p.get("outros_enderecos") or [])
    sede = p.get("sede") or {}
    anos = g["spe_por_ano"]
    aviso = "" if fonte_p == "v2" else ('<p class="aviso">Pesquisa profunda ainda não concluída para esta construtora; a ficha mostra a primeira rodada e a base da Receita.</p>' if fonte_p else
                                        '<p class="aviso">Sem pesquisa web; a ficha mostra só a base da Receita.</p>')
    return f"""
<section class="ficha" id="{e(slug)}">
 <header class="topo">
  <div class="linha"><span class="pos">#{pos}</span><span class="seg">{e(g['segmento'])}</span>{'<span class="seg alerta">gigante</span>' if g['gigante'] else ''}</div>
  <h2>{e(marca)}</h2>
  <p class="sub">{link(site)} · CNPJ {fmt_cnpj(g_(p, 'cnpj_principal') or g['cnpj_sede'])} · {e(g_(p, 'razao_social_principal') or g['razao_sede'])}</p>
  <div class="kpis">
   <div><b>{e(ad.get('nota_0a10', '–'))}</b><span>aderência Hadron (0–10)</span></div>
   <div><b>{e(p.get('confianca_0a10', '–'))}</b><span>confiança dos dados</span></div>
   <div><b>{pct(v)}%</b><span>validadores</span></div>
   <div><b>{len(web)}</b><span>empreendimentos na web</span></div>
   <div><b>{g['n_spe']}</b><span>SPEs na Receita</span></div>
   <div><b>{len(em)}</b><span>e-mails</span></div>
  </div>
 </header>
 {aviso}
 <div class="abordar">
  <h3>Como abordar</h3>
  <p>{e(ad.get('como_abordar') or '—')}</p>
  <p><b>Canal para oferecer projeto:</b> {e(canal.get('descricao') or '—')} {link(canal.get('url_ou_contato')) if str(canal.get('url_ou_contato', '')).startswith('http') else e(canal.get('url_ou_contato'))}{fonte(canal.get('fonte'))}</p>
  <p class="muted"><b>Por que está na lista:</b> {e(ad.get('justificativa') or '—')}</p>
 </div>
 {f'<ul class="alertas">{alertas}</ul>' if alertas else ''}

 <h3>Onde atua</h3>
 {mapa(atua)}
 <div class="onde">
  <div>
   <p><b>Cidades:</b> {e(', '.join(at.get('cidades') or []) or '—')}</p>
   <p><b>Bairros de BH:</b> {e(', '.join(at.get('bairros_bh') or g['bairros_top'][:10]) or '—')}</p>
   <p><b>Faixa de preço:</b> {e(at.get('faixa_preco_m2') or '—')} · <b>Porte típico:</b> {e(at.get('porte_tipico') or '—')}</p>
   <p><b>Ritmo:</b> {e(at.get('lancamentos_por_ano') or '—')}</p>
   <p><b>SPEs abertas por ano (Receita):</b> {' · '.join(f'{a}: {n}' for a, n in anos.items()) or '—'}</p>
   <p><b>Bairros-alvo com SPE:</b> {e(', '.join(g['bairros_alvo_medio'] + g['bairros_alvo_leste']) or 'nenhum')}</p>
  </div>
 </div>
 <h4>Empreendimentos encontrados na web ({len(web)})</h4>
 {tabela(['Empreendimento', 'Endereço', 'Região', 'Status', 'Ano', 'Unid.', 'Área', 'Preço / m²', 'Padrão', ''], web_l, 'emp')}
 <p class="nota">* R$/m² calculado a partir de preço e área do mesmo anúncio.</p>
 <h4>SPEs na Receita que não apareceram na web ({len(base)}; o bairro da SPE costuma ser o do terreno)</h4>
 {tabela(['CNPJ', 'Razão social', 'Bairro', 'Regional', 'Aberta em'], base_l, 'spe')}

 <h3>Contatos da empresa</h3>
 <h4>E-mails ({len(em)})</h4>{tabela(['E-mail', 'Área', 'Pessoa', 'Origem', 'Situação', ''], em_l)}
 <h4>Telefones ({len(fo)})</h4>{tabela(['Número', 'Tipo', 'Área', 'Origem', 'Situação', ''], fo_l)}
 <p class="nota">{'Fora da lista: ' + ', '.join(f'{n} {k}' for k, n in desc.items()) + '.' if desc else ''} "Contador" = contato que se repete em 5 ou mais empresas sem relação na Receita.</p>
 <p><b>Sede:</b> {e(sede.get('endereco') or g['bairro_sede'])}{(', ' + e(sede.get('bairro'))) if sede.get('bairro') else ''}{fonte(sede.get('fonte'))}</p>
 {f'<ul>{outros_end}</ul>' if outros_end else ''}
 <p><b>Redes:</b> {' · '.join(x for x in [link(g_(redes, 'instagram', 'url'), 'Instagram') + (f" ({e(g_(redes, 'instagram', 'seguidores'))})" if g_(redes, 'instagram', 'seguidores') else ''),
     link(g_(redes, 'linkedin', 'url'), 'LinkedIn') + (f" ({e(g_(redes, 'linkedin', 'funcionarios'))})" if g_(redes, 'linkedin', 'funcionarios') else ''),
     link(redes.get('facebook'), 'Facebook'), link(redes.get('youtube'), 'YouTube')] if x) or '—'}</p>

 <h3>Quem decide</h3>
 {tabela(['Nome', 'Cargo', 'Papel', 'LinkedIn', 'Histórico', ''], pes_l)}
 <p><b>Controladores finais:</b> {cf or '—'} {('· <b>Holding:</b> ' + e(ctrl.get('holding'))) if ctrl.get('holding') else ''}</p>
 <p><b>Sócios PF na Receita:</b> {e(', '.join(g['socios_pf'][:8]) or '—')}<br><b>Sócios PJ:</b> {e(', '.join(g['socios_pj'][:6]) or '—')}</p>
 <p><b>Parceiros em SPEs:</b> {e(', '.join((ctrl.get('parceiros_frequentes') or []) + g['parceiros'][:5]) or '—')}</p>

 <h3>Parâmetros de validação</h3>
 {tabela(['', 'Validador', 'Resultado', 'Detalhe'], [[f'<span class="ic {s}">{ICONE[s]}</span>', n, ROT[s], d] for n, s, d in v], 'val')}

 <h3>Saúde e reputação</h3>
 <p><b>Situação cadastral:</b> {e(sa.get('situacao_cadastral') or '—')} · <b>Fundação:</b> {e(sa.get('fundacao') or '—')} ·
  <b>Reclame Aqui:</b> {e(' '.join(x for x in (str(ra.get('nota') or ''), str(ra.get('reputacao') or '')) if x) or 'sem página')}{fonte(ra.get('fonte'))} ·
  <b>Recuperação judicial:</b> {e(sa.get('recuperacao_judicial') or '—')}</p>
 {f'<h4>Processos</h4><ul>{proc}</ul>' if proc else ''}
 {f'<h4>Notícias</h4><ul class="noticias">{notic}</ul>' if notic else ''}
 {('<p><b>Prêmios e associações:</b> ' + e(', '.join(sa.get('premios_associacoes') or [])) + '</p>') if sa.get('premios_associacoes') else ''}
 <p class="rodape">Grupo na base: {g['n_cnpj']} CNPJs · capital somado {brl(g['capital_total'])} · posição {g['rank']} no ranking automático.
 Pesquisa em {e(p.get('data_pesquisa') or HOJE)} ({e(p.get('buscas_feitas') or '?')} buscas). Só contatos profissionais publicados pela empresa,
 por órgão oficial ou pela pessoa como representante da empresa.</p>
</section>"""

# ------------------------------------------------------------------ panorama
def panorama():
    ativos = [g for g in G if g["ativo"]]
    anos = collections.Counter(e_["ano"] for g in G for e_ in g["empreendimentos"])
    serie = [(a, anos.get(a, 0)) for a in map(str, range(2015, 2027))]
    reg = collections.Counter(e_["regional"] for g in G for e_ in g["empreendimentos"]
                              if e_["ano"] >= "2021" and not e_["mesmo_bairro_da_sede"] and not e_["regional"].startswith("BH"))
    # quem constroi em cada bairro-alvo (SPEs desde 2021)
    linhas = []
    for lista, tipo in ((BAIRROS_MEDIO, "médio"), (BAIRROS_LESTE, "leste/econômico")):
        for b in lista:
            alvo = SINONIMOS.get(norm(b), {norm(b)})
            cont = collections.Counter(); tot = 0; tot_all = 0
            for g in G:
                n_all = sum(1 for x in g["empreendimentos"] if norm(x["bairro"]) in alvo)
                n = sum(1 for x in g["empreendimentos"] if norm(x["bairro"]) in alvo and x["ano"] >= "2021")
                tot_all += n_all; tot += n
                if n: cont[g["nome"] + (" ★" if g["gid"] in SEL else "")] += n
            linhas.append([f"<b>{e(b)}</b>", tipo, str(tot_all), str(tot),
                           e(", ".join(f"{k} ({v})" for k, v in cont.most_common(6))) or '<span class="muted">nenhuma SPE recente</span>'])
    return ativos, serie, reg, linhas

# ------------------------------------------------------------------ montagem
CSS = """
/* Layout: relatorio de uma coluna; fichas empilhadas; tabelas e graficos rolam no proprio bloco */
:root{--bg:#fafaf8;--fg:#1b2127;--mut:#5d6873;--bd:#dce0e3;--soft:#eff1f0;--ac:#1d6a58;--ac2:#b8ddd2;--ok:#1a7a43;--pa:#9d5d00;--fa:#b3261e;
--warn-bg:#fcf0da;--warn-fg:#6b4300;--display:"Archivo",Arial,sans-serif;--body:"Source Sans 3",-apple-system,Segoe UI,Roboto,sans-serif;--mono:"JetBrains Mono",ui-monospace,Menlo,monospace}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){--bg:#111514;--fg:#e3e8e6;--mut:#98a4a0;--bd:#2a3230;--soft:#191f1d;--ac:#62c3a6;--ac2:#24483e;--ok:#58c98a;--pa:#e8b15b;--fa:#ff8a80;--warn-bg:#34270e;--warn-fg:#f3d6a2;color-scheme:dark}}
:root[data-theme="dark"]{--bg:#111514;--fg:#e3e8e6;--mut:#98a4a0;--bd:#2a3230;--soft:#191f1d;--ac:#62c3a6;--ac2:#24483e;--ok:#58c98a;--pa:#e8b15b;--fa:#ff8a80;--warn-bg:#34270e;--warn-fg:#f3d6a2;color-scheme:dark}
@media print{:root,:root:not([data-theme="light"]){--bg:#fff;--fg:#1b2127;--mut:#5d6873;--bd:#dce0e3;--soft:#eff1f0;--ac:#1d6a58;--ac2:#b8ddd2;--ok:#1a7a43;--pa:#9d5d00;--fa:#b3261e;--warn-bg:#fcf0da;--warn-fg:#6b4300;color-scheme:light}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.5 var(--body)}
main{max-width:1040px;margin:0 auto;padding-block:24px;padding-inline:16px}
a{color:var(--ac);text-decoration:none;overflow-wrap:anywhere}a:hover,a:focus-visible{text-decoration:underline}a.src{font-size:11px;color:var(--mut)}
h1,h2,h3{font-family:var(--display);text-wrap:balance;letter-spacing:-.01em}
h1{font-size:30px;margin:0 0 4px}h2{font-size:24px;margin:4px 0}h3{font-size:17px;margin:26px 0 8px;padding-bottom:4px;border-bottom:2px solid var(--bd)}
h4{font-size:12px;margin:16px 0 6px;color:var(--mut);text-transform:uppercase;letter-spacing:.06em}
table{width:100%;border-collapse:collapse;font-size:13px;font-variant-numeric:tabular-nums}th,td{text-align:left;padding:5px 6px;border-bottom:1px solid var(--bd);vertical-align:top}
th{background:var(--soft);font-weight:600;white-space:nowrap}.wrap{overflow-x:auto;margin-bottom:6px}
.muted,.sub{color:var(--mut)}.vazio,.nota{color:var(--mut);font-size:12px;margin:2px 0 8px}.vazio{font-style:italic}
.mono{font-family:var(--mono);font-size:12px}
.ficha{border-top:3px solid var(--ac);padding-top:14px;margin-top:40px}
.linha{display:flex;gap:8px;align-items:center;flex-wrap:wrap}.pos{font:600 13px var(--mono);color:var(--mut)}
.seg{font-size:11px;padding:2px 8px;border-radius:10px;background:var(--soft);border:1px solid var(--bd)}.seg.alerta{color:var(--fa)}
.kpis{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));gap:8px;margin:12px 0}
.kpis div{background:var(--soft);border-radius:6px;padding:8px 10px;min-width:0}.kpis b{font:600 19px var(--mono);display:block}.kpis span{font-size:11px;color:var(--mut)}
.abordar{background:var(--soft);padding:10px 14px;border-radius:6px}.abordar h3{border:0;margin:0 0 4px;padding:0;color:var(--ac)}.abordar p{margin:4px 0}
.duas{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,1fr);gap:16px;align-items:start}.duas p{margin:4px 0}
.st{font-size:11px;white-space:nowrap;padding:1px 6px;border-radius:8px;border:1px solid var(--bd)}
.st.confirmado{color:var(--ok)}.st.publicado{color:var(--pa)}.st.contador,.st.domínio_de_terceiro{color:var(--fa)}.st.cadastro_Receita{color:var(--mut)}
.ic{font-weight:700}.ic.ok{color:var(--ok)}.ic.parcial{color:var(--pa)}.ic.falha{color:var(--fa)}.ic.nd{color:var(--mut)}table.val td:first-child{width:22px;text-align:center}
.tom{font-size:11px;padding:0 5px;border-radius:6px;background:var(--soft)}.tom.negativo{color:var(--fa)}.tom.positivo{color:var(--ok)}
.alertas,.aviso{background:var(--warn-bg);color:var(--warn-fg);border-radius:6px;padding:8px 8px 8px 26px}.aviso{padding:8px 12px}
.chart{width:100%;height:auto;max-width:100%;display:block}.chart .grid{stroke:var(--bd);stroke-width:1}.chart .ax{fill:var(--mut);font:11px var(--body)}
.chart .lbl{fill:var(--fg);font:600 11px var(--mono)}.chart .bar path,.chart .bar rect:not(.hit){fill:var(--ac)}.chart .hit{fill:transparent}
.chart .bar:hover path,.chart .bar:hover rect:not(.hit){opacity:.8}
.mapa .bh{fill:var(--soft);stroke:var(--bd);stroke-width:1.5}.mapa .pt circle{fill:var(--bd)}.mapa .pt.on circle{fill:var(--ac);fill-opacity:.55;stroke:var(--bg);stroke-width:2}
.mapas{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,1fr);gap:12px}.onde p{margin:4px 0}
.grade{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,1fr);gap:20px}
.stats{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:8px;margin:12px 0}.stats div{border-left:3px solid var(--ac);padding:2px 10px}
.stats b{font:600 22px var(--mono);display:block}.stats span{font-size:12px;color:var(--mut)}
.filtro{display:flex;gap:8px;flex-wrap:wrap;margin:8px 0}.filtro input,.filtro select{font:14px var(--body);padding:6px 8px;border:1px solid var(--bd);border-radius:6px;background:var(--bg);color:var(--fg)}
.filtro input:focus-visible,.filtro select:focus-visible{outline:2px solid var(--ac)}
.rodape{font-size:11px;color:var(--mut);margin-top:14px}.voltar{margin:0}
@media (max-width:760px){.mapas{grid-template-columns:1fr}.kpis{grid-template-columns:repeat(3,minmax(0,1fr))}.duas,.grade{grid-template-columns:1fr}.stats{grid-template-columns:repeat(2,minmax(0,1fr))}}
@page{size:A4;margin:11mm}
@media print{.ficha{break-before:page;margin-top:0;border-top:0}h3,h4{break-after:avoid}tr,svg{break-inside:avoid}a.src{display:none}main{padding:0}.voltar,.filtro{display:none}}
"""
FONTES = ('<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Archivo:wght@600;700'
          '&family=JetBrains+Mono:wght@400;600&family=Source+Sans+3:wght@400;600;700&display=swap">')

def pagina(titulo, corpo, esqueleto=True, js=""):
    miolo = f"<title>{e(titulo)}</title>{FONTES}<style>{CSS}</style><main>{corpo}</main>{js}"
    if not esqueleto: return miolo
    return ('<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">'
            + miolo.replace("<main>", "</head><body><main>", 1) + "</body></html>")

JS_FILTRO = """<script>
(function(){var q=document.getElementById('q'),s=document.getElementById('seg'),r=document.getElementById('reg'),t=document.getElementById('tab-op');
if(!t)return;var rows=[].slice.call(t.tBodies[0].rows),n=document.getElementById('n-op');
function f(){var a=(q.value||'').toLowerCase(),b=s.value,c=r.value,k=0;rows.forEach(function(tr){var ok=(!a||tr.textContent.toLowerCase().indexOf(a)>=0)&&(!b||tr.dataset.seg===b)&&(!c||(tr.dataset.reg||'').indexOf(c)>=0);tr.hidden=!ok;if(ok)k++});n.textContent=k}
[q,s,r].forEach(function(el){el.addEventListener('input',f)});f()})();
</script>"""

def main():
    itens = []
    for gid, slug in zip(SEL, SLUGS):
        g = GID[gid]; p, fonte_p = carregar(slug, g)
        em, fo, desc = contatos(g, p)
        web, base = empreendimentos(g, p)
        v = validadores(g, p, em, fo, web)
        ad = num(g_(p, "aderencia", "nota_0a10"), 0); cf = num(p.get("confianca_0a10"), 0)
        prior = round(5 * ad + 2.5 * cf + 0.25 * g["score"], 1)
        itens.append(dict(g=g, p=p, slug=slug, fonte=fonte_p, em=em, fo=fo, desc=desc, web=web, base=base, v=v,
                          atua=onde_atua(web, base), prior=prior))
    itens.sort(key=lambda t: -t["prior"])

    ativos, serie, reg, bairro_linhas = panorama()
    tot_web = sum(len(i["web"]) for i in itens); tot_em = sum(len(i["em"]) for i in itens); tot_pes = sum(len(i["p"].get("pessoas") or []) for i in itens)
    atua_total = collections.Counter()
    for i in itens: atua_total.update(i["atua"])

    rank_l = []
    for n, i in enumerate(itens, 1):
        g, p = i["g"], i["p"]
        pes = p.get("pessoas") or []
        dec = next((x for x in pes if x.get("papel") in ("incorporacao", "novos_negocios", "terrenos", "produto", "projetos")), None) or (pes[0] if pes else {})
        em_best = next((x["email"] for x in i["em"] if x["status"] == "confirmado"), "")
        fo_best = next((x["numero"] for x in i["fo"] if x["status"] == "confirmado"), "")
        onde = ", ".join(k for k, _ in i["atua"].most_common(3))
        rank_l.append([str(n), f'<a href="{e(i["slug"])}.html"><b>{e(g_(p, "marca") or g["nome"])}</b></a><br><span class="muted">{e(g_(p, "site", "url") or g["dominio"])}</span>',
                       e(g["segmento"]), f'<span class="mono">{i["prior"]}</span>', e(g_(p, "aderencia", "nota_0a10", default="–")),
                       f'{pct(i["v"])}%', str(len(i["web"])), e(onde),
                       e(dec.get("nome") or "—") + (f'<br><span class="muted">{e(dec.get("cargo"))}</span>' if dec.get("cargo") else ""),
                       f'<span class="mono">{e(em_best)}</span><br><span class="mono">{e(fo_best)}</span>'])

    intro = f"""
<h1>Construtoras-alvo · BH e RMBH</h1>
<p class="sub">Mapeamento para apresentar projetos novos · {HOJE} · base da Receita (jul/2026) + pesquisa web</p>
<div class="stats">
 <div><b>{len(G):,}</b><span>grupos com incorporação na base</span></div>
 <div><b>{len(ativos):,}</b><span>ativos (SPE aberta desde 2021)</span></div>
 <div><b>{len(itens)}</b><span>construtoras com ficha completa</span></div>
 <div><b>{tot_web}</b><span>empreendimentos localizados na web</span></div>
</div>
<h3>Critérios (conversa do grupo Hadron – Estratégico)</h3>
<p>Prédios com mais de 36 unidades · venda de R$ 10–14 mil/m² (sem alto padrão) ou MCMV em BH, Contagem, Betim e Ribeirão das Neves ·
empresa com site · área construída acima de 3.000 m² · bairros sugeridos: {', '.join(BAIRROS_MEDIO)}; e {', '.join(BAIRROS_LESTE)}.</p>
<h3>Como a base foi percorrida</h3>
<p>Os 15.681 CNPJs de construção e incorporação de BH foram agrupados em grupos econômicos: mesmo site, mesma raiz de CNPJ,
mesmo sócio pessoa física e holding. Um CNPJ sem site entra no grupo do site que predomina entre os sócios dele, e SPE em parceria não funde duas
construtoras. Cada SPE virou uma pista de empreendimento, com bairro e regional. Contatos que se repetem em 5 ou mais empresas sem relação foram
marcados como do contador. A nota automática pesa volume e ritmo de SPEs, a fatia de SPEs nos bairros-alvo, o segmento, o site e os contatos
próprios. Gigantes com projeto interno (MRV, Direcional, Cyrela, Patrimar, Emccamp, Riva, Conata, VIC) foram rebaixados.
Das oportunidades ativas, {len(itens)} construtoras receberam pesquisa web profunda: portfólio com local, todos os e-mails e telefones publicados,
decisores com LinkedIn, controle, reputação e forma de abordagem. No total, {tot_em} e-mails e {tot_pes} pessoas.</p>
<p><b>Prioridade</b> = 5 × aderência (0–10) + 2,5 × confiança (0–10) + ¼ da nota automática da base.</p>

<h3>Mercado na base da Receita</h3>
<div class="grade">
 <div><h4>SPEs abertas por ano (todas as construtoras de BH)</h4>{barras(serie)}</div>
 <div><h4>SPEs abertas desde 2021, por regional (endereço da SPE)</h4>{barras_h(reg.most_common())}</div>
</div>
<h3>Quem está construindo nos bairros-alvo</h3>
<p class="nota">SPEs com endereço no bairro. "Desde 2021" indica atividade recente. ★ = construtora com ficha neste relatório.</p>
{tabela(['Bairro', 'Lista', 'SPEs (total)', 'Desde 2021', 'Construtoras com SPE desde 2021 (nº de SPEs)'], bairro_linhas)}

<h3>Onde as {len(itens)} construtoras atuam</h3>
<p class="nota">Empreendimentos localizados na web e SPEs da Receita (sem contar SPE registrada no endereço da sede).</p>
{mapa(atua_total)}

<h3>Ranking das construtoras com ficha</h3>
{tabela(['#', 'Construtora', 'Segmento', 'Prioridade', 'Aderência', 'Validação', 'Empreend.', 'Onde mais atua', 'Decisor', 'E-mail · telefone confirmados'], rank_l, 'rank')}
<p><a href="oportunidades.html"><b>Lista completa das {len([g for g in G if g['ativo']])} oportunidades ativas da base →</b></a></p>

<h3>Limites</h3>
<ul><li>O ambiente não abre sites diretamente; a pesquisa usou o buscador. Contato marcado "publicado?" apareceu em trecho de busca e vale confirmar.</li>
<li>A base da Receita é só de CNPJs com sede em BH. Obras na RMBH vêm da pesquisa web.</li>
<li>O bairro da SPE nem sempre é o do terreno. Quando coincide com o da sede, fica marcado "(= sede)".</li>
<li>Não entram celulares nem e-mails genéricos (gmail etc.) da Receita, que podem ser pessoais; cada ficha informa quantos ficaram de fora.</li>
<li>A regional de cada bairro e as posições do mapa são aproximadas.</li></ul>"""

    # lista completa de oportunidades
    op = [g for g in G if g["ativo"]]
    sel_slug = {gid: s for gid, s in zip(SEL, SLUGS)}
    op_l = []
    for g in op:
        regs = " ".join(g["regionais"])
        nome = f'<a href="{e(sel_slug[g["gid"]])}.html"><b>{e(g["nome"])}</b></a> ★' if g["gid"] in sel_slug else f'<b>{e(g["nome"])}</b>'
        op_l.append((g["segmento"], regs, [str(g["rank"]), nome + f'<br><span class="muted">{e(g["dominio"] or "sem site")}</span>', e(g["segmento"]),
                     f'<span class="mono">{g["score"]}</span>', str(g["n_spe"]), str(g["spe_desde_2023"]),
                     e(", ".join(f"{k} {v}" for k, v in list(g["regionais"].items())[:3])), e(", ".join(g["bairros_top"][:4])),
                     "<br>".join(f'<span class="mono">{e(x["email"])}</span>' for x in g["emails"] if x["classe"] == "dominio_proprio")[:400],
                     "<br>".join(f'<span class="mono">{e(x["numero"])}</span>' for x in g["fones"] if x["classe"] == "proprio_receita" and x["tipo"] == "fixo")[:200],
                     e(", ".join(g["socios_pf"][:3]))]))
    cab = ["#", "Grupo", "Segmento", "Nota", "SPEs", "desde 2023", "Regionais", "Bairros", "E-mails do domínio", "Fixos (Receita)", "Sócios"]
    trs = "".join(f'<tr data-seg="{e(s)}" data-reg="{e(r)}">' + "".join(f"<td>{c}</td>" for c in l) + "</tr>" for s, r, l in op_l)
    opts_seg = "".join(f'<option>{e(s)}</option>' for s in ["Médio", "MCMV/econômico", "Misto", "Alto padrão"])
    opts_reg = "".join(f'<option>{e(r)}</option>' for r in REGIONAIS)
    op_html = f"""<p class="voltar"><a href="index.html">← voltar</a></p>
<h1>Oportunidades ativas na base</h1>
<p class="sub">{len(op)} grupos com SPE aberta desde 2021, ordenados pela nota automática. ★ = tem ficha completa.
Contatos aqui são os da Receita (e-mail no domínio da empresa e telefone fixo que não se repete em outras empresas).</p>
<div class="filtro"><label for="q" class="muted">Buscar</label><input id="q" type="search" placeholder="nome, bairro, sócio…">
<label for="seg" class="muted">Segmento</label><select id="seg"><option value="">todos</option>{opts_seg}</select>
<label for="reg" class="muted">Regional</label><select id="reg"><option value="">todas</option>{opts_reg}</select>
<span class="muted"><b id="n-op">{len(op)}</b> grupos</span></div>
<div class="wrap"><table id="tab-op"><thead><tr>{''.join(f'<th>{c}</th>' for c in cab)}</tr></thead><tbody>{trs}</tbody></table></div>"""

    OUT.mkdir(exist_ok=True)
    fichas = [ficha(i, n) for n, i in enumerate(itens, 1)]
    for i, f in zip(itens, fichas):
        (OUT / f"{i['slug']}.html").write_text(pagina(g_(i["p"], "marca") or i["g"]["nome"], '<p class="voltar"><a href="index.html">← voltar ao ranking</a></p>' + f), encoding="utf-8")
    (OUT / "index.html").write_text(pagina("Construtoras-alvo BH e RMBH", intro), encoding="utf-8")
    (OUT / "artifact_index.html").write_text(pagina("Construtoras-alvo BH e RMBH", intro, esqueleto=False), encoding="utf-8")
    (OUT / "oportunidades.html").write_text(pagina("Oportunidades ativas", op_html, js=JS_FILTRO), encoding="utf-8")
    top_op = tabela(cab[:8], [l[:8] for _, _, l in op_l[:200]])
    (OUT / "relatorio.html").write_text(pagina("Construtoras-alvo BH e RMBH", intro + "".join(fichas) +
        f'<section class="ficha"><h2>Anexo · 200 primeiras oportunidades ativas da base</h2>{top_op}</section>'), encoding="utf-8")
    planilha(itens, op)
    print(f"{len(itens)} fichas ({sum(i['fonte'] == 'v2' for i in itens)} v2, {sum(i['fonte'] == 'v1' for i in itens)} v1) · "
          f"{tot_web} empreendimentos web · {tot_em} e-mails · {tot_pes} pessoas · {len(op)} oportunidades")

def planilha(itens, op):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    def aba(nome, cab, linhas, larg=None):
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1D6A58")
        for l in linhas: ws.append([("" if x is None else x) for x in l])
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for i, _ in enumerate(cab):
            col = ws.cell(1, i + 1).column_letter
            ws.column_dimensions[col].width = (larg or {}).get(i, min(45, max(10, max((len(str(r[i])) for r in [cab] + linhas[:300] if i < len(r)), default=10) + 2)))
        return ws
    wb.remove(wb.active)
    aba("Ranking", ["#", "Construtora", "Site", "Segmento", "Prioridade", "Aderência", "Confiança", "Validação %", "Empreend. web", "SPEs Receita",
                    "Onde mais atua", "Como abordar", "Canal para projeto"],
        [[n, g_(i["p"], "marca") or i["g"]["nome"], g_(i["p"], "site", "url") or i["g"]["dominio"], i["g"]["segmento"], i["prior"],
          g_(i["p"], "aderencia", "nota_0a10"), i["p"].get("confianca_0a10"), pct(i["v"]), len(i["web"]), i["g"]["n_spe"],
          ", ".join(f"{k} ({v})" for k, v in i["atua"].most_common(4)), g_(i["p"], "aderencia", "como_abordar"),
          g_(i["p"], "canal_para_oferecer_projeto", "descricao")] for n, i in enumerate(itens, 1)], {11: 60, 12: 40})
    aba("E-mails", ["Construtora", "E-mail", "Área", "Pessoa", "Origem", "Situação", "Fonte"],
        [[g_(i["p"], "marca") or i["g"]["nome"], x["email"], x["area"], x["pessoa"], x["origem"], x["status"], x["fonte"] or ""] for i in itens for x in i["em"]])
    aba("Telefones", ["Construtora", "Número", "Tipo", "Área", "Origem", "Situação", "Fonte"],
        [[g_(i["p"], "marca") or i["g"]["nome"], x["numero"], x["tipo"], x["area"], x["origem"], x["status"], x["fonte"] or ""] for i in itens for x in i["fo"]])
    aba("Pessoas", ["Construtora", "Nome", "Cargo", "Papel", "LinkedIn", "Histórico", "Fonte"],
        [[g_(i["p"], "marca") or i["g"]["nome"], x.get("nome"), x.get("cargo"), x.get("papel"), x.get("linkedin"), x.get("formacao_ou_historico"), x.get("fonte")]
         for i in itens for x in i["p"].get("pessoas") or []])
    aba("Empreendimentos web", ["Construtora", "Empreendimento", "Endereço", "Bairro", "Cidade", "Região", "Status", "Ano", "Unidades", "Área privativa",
                                "Preço", "Preço/m²", "Calculado", "Padrão", "MCMV/padrão", "Fonte"],
        [[g_(i["p"], "marca") or i["g"]["nome"], x.get("nome"), x.get("endereco"), x.get("bairro"), x.get("cidade"), x.get("onde"), x.get("status"), x.get("ano"),
          x.get("unidades"), x.get("area_privativa"), x.get("preco"), x.get("preco_m2"), "sim" if x.get("calculado") else "", x.get("padrao"), x.get("tipologia"), x.get("fonte")]
         for i in itens for x in i["web"]])
    aba("SPEs Receita (todas)", ["Grupo", "Rank grupo", "CNPJ", "Razão social", "Bairro", "Regional", "Aberta em", "Capital", "Bairro = sede"],
        [[g["nome"], g["rank"], fmt_cnpj(x["cnpj"]), x["razao"], x["bairro"], x["regional"], x["ano"], x["capital"], "sim" if x["mesmo_bairro_da_sede"] else ""]
         for g in G if g["ativo"] for x in g["empreendimentos"]])
    aba("Oportunidades", ["Rank", "Grupo", "Site", "Segmento", "Nota", "SPEs", "SPEs desde 2023", "Regionais", "Bairros", "E-mails do domínio",
                          "Fixos (Receita)", "Sócios PF", "Sócios PJ", "CNPJ sede", "Bairro sede", "Gigante"],
        [[g["rank"], g["nome"], g["dominio"], g["segmento"], g["score"], g["n_spe"], g["spe_desde_2023"],
          ", ".join(f"{k} {v}" for k, v in g["regionais"].items()), ", ".join(g["bairros_top"][:6]),
          ", ".join(x["email"] for x in g["emails"] if x["classe"] == "dominio_proprio"),
          ", ".join(x["numero"] for x in g["fones"] if x["classe"] == "proprio_receita" and x["tipo"] == "fixo"),
          ", ".join(g["socios_pf"][:5]), ", ".join(g["socios_pj"][:3]), fmt_cnpj(g["cnpj_sede"]), g["bairro_sede"], "sim" if g["gigante"] else ""] for g in op])
    wb.save(R / "Construtoras_RMBH_mapeamento.xlsx")

if __name__ == "__main__":
    main()
